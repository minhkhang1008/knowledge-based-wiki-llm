import os
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from datetime import datetime

from app.services.document_parser.table_extractor import PptxTableExtractor
from app.services.document_parser.image_extractor import PptxImageExtractor
from app.services.document_parser.spatial_analyzer import PptxSpatialAnalyzer
from app.services.document_parser.converter_registry import register

class PptxCoreParser:
    def __init__(self, file_path: str, images_output_dir: str = None):
        self.file_path = file_path
        self.prs = Presentation(file_path)
        self.images_output_dir = images_output_dir or os.path.join(
            os.path.dirname(file_path), "extracted_images"
        )
        
    def _generate_front_matter(self) -> str:
        filename = os.path.basename(self.file_path)
        num_slides = len(self.prs.slides)
        date_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        return f"---\ntitle: {filename}\ndate_extracted: {date_str}\ntotal_slides: {num_slides}\n---\n\n"

    def _extract_slide_title(self, slide) -> str:
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            return f"## {slide.shapes.title.text.strip()}\n\n"
        return "## [Untitled Slide]\n\n"

    def _extract_text_frame(self, text_frame) -> str:
        content = ""
        for paragraph in text_frame.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            indent_level = paragraph.level if paragraph.level else 0
            indent_space = "    " * indent_level
            content += f"{indent_space}* {text}\n"
        return content + "\n"

    def parse(self) -> str:
        markdown_output = self._generate_front_matter()
        
        for i, slide in enumerate(self.prs.slides):
            markdown_output += self._extract_slide_title(slide)
            
            sorted_shapes = PptxSpatialAnalyzer.sort_shapes(slide.shapes)
            for j, shape in enumerate(sorted_shapes):
                if shape == slide.shapes.title:
                    continue
                if shape.has_table:
                    markdown_output += PptxTableExtractor.extract(shape)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    markdown_output += PptxImageExtractor.extract(
                        shape, self.images_output_dir, i + 1, j + 1
                    )
                elif shape.has_text_frame:
                    markdown_output += self._extract_text_frame(shape.text_frame)
            
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    markdown_output += "> **Speaker Notes:**\n"
                    for line in notes_text.split("\n"):
                        markdown_output += f"> {line}\n"
                    markdown_output += "\n"
                    
            markdown_output += "---\n\n"

        return markdown_output

@register(".pptx")
def pptx_to_markdown(file_path: str | Path, output_dir: str) -> str:
    images_dir = os.path.join(output_dir, "images")
    parser = PptxCoreParser(str(file_path), images_output_dir=images_dir)
    
    # BẮT BUỘC PHẢI CÓ CHỮ 'return' Ở ĐÂY
    return parser.parse()